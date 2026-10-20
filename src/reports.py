"""
Модуль аналитических отчетов и экспорта документов для СУБД (Лабораторная работа №3).
НИУ «МЭИ», Кафедра УИТ, Дисциплина «СУБД», Вариант №7.
"""

import os
import sqlite3
from datetime import datetime
from pathlib import Path


def _format_money(val: int) -> str:
    """Форматирование числа с разделением разрядов пробелом."""
    try:
        return f"{int(val):,}".replace(",", " ")
    except Exception:
        return str(val)


def build_contest_report(db_path: Path, filter_where: str = ""):
    """
    Форма 1: «Распределение НИР по конкурсам грантов»
    Строка = один конкурс.
    Колонки: № п/п, Код конкурса, Название конкурса, Количество НИР, План финансирования, Количество вузов (DISTINCT).
    Строка ИТОГО: сумма НИР, сумма плана, ОБЩЕЕ число уникальных вузов во всей выборке.
    """
    conn = sqlite3.connect(db_path)
    cur = conn.cursor()

    where_sql = f"WHERE {filter_where}" if filter_where else ""

    query = f"""
    SELECT 
        k.codkon,
        k.konk_name,
        COUNT(p.id) AS nir_count,
        COALESCE(SUM(p.plan_fin), 0) AS total_plan,
        COUNT(DISTINCT p.codvuz) AS vuz_count
    FROM gr_konk k
    JOIN gr_proj p ON k.codkon = p.codkon
    {where_sql}
    GROUP BY k.codkon, k.konk_name
    ORDER BY k.codkon
    """
    cur.execute(query)
    raw_rows = cur.fetchall()

    headers = [
        "№ п/п",
        "Код конкурса",
        "Наименование конкурса",
        "Количество НИР",
        "План финансирования",
        "Количество вузов",
    ]

    rows = []
    total_nir = 0
    total_plan = 0
    for idx, r in enumerate(raw_rows, 1):
        codkon, name, n_cnt, p_sum, v_cnt = r
        total_nir += n_cnt
        total_plan += p_sum
        rows.append([
            idx,
            codkon,
            name,
            n_cnt,
            p_sum,
            v_cnt,
        ])

    # Подсчет общего количества уникальных вузов во всей отфильтрованной выборке
    tot_query = f"SELECT COUNT(DISTINCT codvuz) FROM gr_proj {where_sql}"
    cur.execute(tot_query)
    total_distinct_vuz = cur.fetchone()[0]
    conn.close()

    totals = [
        "ИТОГО",
        "",
        "",
        total_nir,
        total_plan,
        total_distinct_vuz,
    ]

    return {
        "title": "Распределение проектов НИР по конкурсам грантов",
        "headers": headers,
        "rows": rows,
        "totals": totals,
        "numeric_cols": [3, 4, 5],
        "money_cols": [4],
    }


def build_vuz_report(db_path: Path, filter_where: str = ""):
    """
    Форма 2: «Распределение НИР по вузам»
    Строка = один вуз.
    Колонки: № п/п, Код вуза, Вуз (аббревиатура), Город, Субъект РФ, Количество НИР, План финансирования, Количество конкурсов (DISTINCT).
    Строка ИТОГО: сумма НИР, сумма плана, ОБЩЕЕ число уникальных конкурсов во всей выборке.
    """
    conn = sqlite3.connect(db_path)
    cur = conn.cursor()

    where_sql = f"WHERE {filter_where}" if filter_where else ""

    query = f"""
    SELECT 
        v.codvuz,
        COALESCE(v.vuz_short_name, v.vuz_name) AS vuz_name,
        COALESCE(v.city, '') AS city,
        COALESCE(v.obl_name, '') AS obl_name,
        COUNT(p.id) AS nir_count,
        COALESCE(SUM(p.plan_fin), 0) AS total_plan,
        COUNT(DISTINCT p.codkon) AS contest_count
    FROM gr_proj p
    JOIN vuz v ON p.codvuz = v.codvuz
    {where_sql}
    GROUP BY v.codvuz, v.vuz_short_name, v.city, v.obl_name
    ORDER BY nir_count DESC, total_plan DESC, v.codvuz ASC
    """
    cur.execute(query)
    raw_rows = cur.fetchall()

    headers = [
        "№ п/п",
        "Код вуза",
        "Вуз",
        "Город",
        "Субъект РФ",
        "Количество НИР",
        "План финансирования",
        "Количество конкурсов",
    ]

    rows = []
    total_nir = 0
    total_plan = 0
    for idx, r in enumerate(raw_rows, 1):
        codvuz, name, city, obl, n_cnt, p_sum, c_cnt = r
        total_nir += n_cnt
        total_plan += p_sum
        rows.append([
            idx,
            codvuz,
            name,
            city,
            obl,
            n_cnt,
            p_sum,
            c_cnt,
        ])

    tot_query = f"SELECT COUNT(DISTINCT codkon) FROM gr_proj {where_sql}"
    cur.execute(tot_query)
    total_distinct_konk = cur.fetchone()[0]
    conn.close()

    totals = [
        "ИТОГО",
        "",
        "",
        "",
        "",
        total_nir,
        total_plan,
        total_distinct_konk,
    ]

    return {
        "title": "Распределение проектов НИР по вузам-исполнителям",
        "headers": headers,
        "rows": rows,
        "totals": totals,
        "numeric_cols": [5, 6, 7],
        "money_cols": [6],
    }


def build_region_report(db_path: Path, filter_where: str = ""):
    """
    Форма 3: «Распределение НИР по регионам (субъектам РФ)»
    Строка = один субъект РФ.
    Колонки: № п/п, Код региона, Субъект РФ, Федеральный округ, Количество НИР, План финансирования, Количество вузов (DISTINCT).
    Строка ИТОГО: сумма НИР, сумма плана, ОБЩЕЕ число уникальных вузов в выборке.
    """
    conn = sqlite3.connect(db_path)
    cur = conn.cursor()

    where_sql = f"WHERE {filter_where}" if filter_where else ""

    query = f"""
    SELECT 
        COALESCE(v.obl_code, '') AS obl_code,
        COALESCE(v.obl_name, 'Не указан') AS obl_name,
        COALESCE(v.region, '') AS region,
        COUNT(p.id) AS nir_count,
        COALESCE(SUM(p.plan_fin), 0) AS total_plan,
        COUNT(DISTINCT p.codvuz) AS vuz_count
    FROM gr_proj p
    JOIN vuz v ON p.codvuz = v.codvuz
    {where_sql}
    GROUP BY v.obl_code, v.obl_name, v.region
    ORDER BY nir_count DESC, total_plan DESC, v.obl_name ASC
    """
    cur.execute(query)
    raw_rows = cur.fetchall()

    headers = [
        "№ п/п",
        "Код региона",
        "Субъект РФ",
        "Федеральный округ",
        "Количество НИР",
        "План финансирования",
        "Количество вузов",
    ]

    rows = []
    total_nir = 0
    total_plan = 0
    for idx, r in enumerate(raw_rows, 1):
        obl_code, obl_name, reg, n_cnt, p_sum, v_cnt = r
        total_nir += n_cnt
        total_plan += p_sum
        rows.append([
            idx,
            obl_code,
            obl_name,
            reg,
            n_cnt,
            p_sum,
            v_cnt,
        ])

    tot_query = f"SELECT COUNT(DISTINCT codvuz) FROM gr_proj {where_sql}"
    cur.execute(tot_query)
    total_distinct_vuz = cur.fetchone()[0]
    conn.close()

    totals = [
        "ИТОГО",
        "",
        "",
        "",
        total_nir,
        total_plan,
        total_distinct_vuz,
    ]

    return {
        "title": "Распределение проектов НИР по субъектам РФ",
        "headers": headers,
        "rows": rows,
        "totals": totals,
        "numeric_cols": [4, 5, 6],
        "money_cols": [5],
    }


def export_report_to_file(file_path: str, report_data: dict, filter_summary: str = ""):
    """
    Экспорт отчета в выбранный файл (HTML, XLSX, CSV или TXT).
    Строго включает:
    1. Общий заголовок отчета;
    2. Подзаголовок с параметрами фильтрации (ОБЯЗАТЕЛЬНО по ТЗ Полотнова);
    3. Дату и время формирования;
    4. Таблицу с данными и строкой «ИТОГО».
    """
    ext = Path(file_path).suffix.lower()
    Path(file_path).parent.mkdir(parents=True, exist_ok=True)
    title = report_data.get("title", "Отчет")
    headers = report_data.get("headers", [])
    rows = report_data.get("rows", [])
    totals = report_data.get("totals", [])
    money_cols = set(report_data.get("money_cols", []))
    numeric_cols = set(report_data.get("numeric_cols", []))

    filter_desc = filter_summary if filter_summary else "не установлены (полная база данных)"
    now_str = datetime.now().strftime("%d.%m.%Y %H:%M")

    if ext == ".xlsx":
        _export_to_excel(file_path, title, filter_desc, now_str, headers, rows, totals, money_cols, numeric_cols)
    elif ext in (".html", ".htm"):
        _export_to_html(file_path, title, filter_desc, now_str, headers, rows, totals, money_cols, numeric_cols)
    elif ext == ".csv":
        _export_to_csv(file_path, title, filter_desc, now_str, headers, rows, totals)
    else:  # По умолчанию .txt
        _export_to_txt(file_path, title, filter_desc, now_str, headers, rows, totals, money_cols)


def _export_to_excel(file_path, title, filter_desc, now_str, headers, rows, totals, money_cols, numeric_cols):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчет"
    ws.views.sheetView[0].showGridLines = True

    # 1. Заголовок
    ws.cell(row=1, column=1, value=title).font = Font(name="Calibri", size=14, bold=True)
    ws.cell(row=2, column=1, value=f"Параметры фильтрации: {filter_desc}").font = Font(name="Calibri", size=11, italic=True)
    ws.cell(row=3, column=1, value=f"Дата формирования: {now_str}").font = Font(name="Calibri", size=9, color="555555")

    # 2. Шапка таблицы
    thin_border = Border(
        left=Side(style='thin', color='B0B0B0'),
        right=Side(style='thin', color='B0B0B0'),
        top=Side(style='thin', color='B0B0B0'),
        bottom=Side(style='thin', color='B0B0B0')
    )
    header_fill = PatternFill(start_color="E6EEF8", end_color="E6EEF8", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True)

    header_row_idx = 5
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 3. Данные
    data_font = Font(name="Calibri", size=11)
    current_row = header_row_idx + 1
    for r in rows:
        for col_idx, val in enumerate(r, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if (col_idx - 1) in money_cols:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif (col_idx - 1) in numeric_cols:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        current_row += 1

    # 4. Строка ИТОГО
    if totals:
        total_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        total_font = Font(name="Calibri", size=11, bold=True)
        thick_top_border = Border(
            left=Side(style='thin', color='B0B0B0'),
            right=Side(style='thin', color='B0B0B0'),
            top=Side(style='medium', color='333333'),
            bottom=Side(style='double', color='333333')
        )
        for col_idx, val in enumerate(totals, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = total_font
            cell.fill = total_fill
            cell.border = thick_top_border
            if (col_idx - 1) in money_cols:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif (col_idx - 1) in numeric_cols:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Автоподбор ширины столбцов
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row >= header_row_idx and cell.value:
                val_str = _format_money(cell.value) if isinstance(cell.value, (int, float)) else str(cell.value)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(file_path)


def _export_to_html(file_path, title, filter_desc, now_str, headers, rows, totals, money_cols, numeric_cols):
    html = f"""<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Arial, sans-serif;
    margin: 30px;
    color: #222;
    background: #fff;
  }}
  h1 {{
    font-size: 20px;
    margin-bottom: 6px;
    color: #1a365d;
  }}
  .meta {{
    font-size: 13px;
    color: #4a5568;
    margin-bottom: 4px;
  }}
  .filter-box {{
    background: #edf2f7;
    border-left: 4px solid #3182ce;
    padding: 8px 12px;
    margin: 12px 0 20px 0;
    font-size: 13px;
    font-weight: 500;
  }}
  table {{
    border-collapse: collapse;
    width: 100%;
    margin-top: 10px;
    font-size: 13px;
  }}
  th, td {{
    border: 1px solid #cbd5e0;
    padding: 7px 10px;
  }}
  th {{
    background: #ebf8ff;
    color: #2b6cb0;
    font-weight: 600;
    text-align: center;
  }}
  tr:nth-child(even) td {{
    background: #f7fafc;
  }}
  tr:hover td {{
    background: #edf2f7;
  }}
  .text-left {{ text-align: left; }}
  .text-center {{ text-align: center; }}
  .text-right {{ text-align: right; font-variant-numeric: tabular-nums; }}
  .total-row td {{
    background: #fffaf0 !important;
    font-weight: bold;
    border-top: 2px solid #4a5568;
    border-bottom: 2px solid #4a5568;
  }}
  @media print {{
    body {{ margin: 10mm; }}
    .filter-box {{ background: none; border-left: 2px solid #000; }}
  }}
</style>
</head>
<body>
  <h1>{title}</h1>
  <div class="meta">Дата формирования: {now_str} | СУБД: Сопровождение конкурсов на соискание грантов (Вариант 7)</div>
  <div class="filter-box">Параметры фильтрации: {filter_desc}</div>
  
  <table>
    <thead>
      <tr>
"""
    for h in headers:
        html += f"        <th>{h}</th>\n"
    html += "      </tr>\n    </thead>\n    <tbody>\n"

    for r in rows:
        html += "      <tr>\n"
        for col_idx, val in enumerate(r):
            if col_idx in money_cols:
                cls = "text-right"
                val_str = _format_money(val)
            elif col_idx in numeric_cols:
                cls = "text-right"
                val_str = str(val)
            elif col_idx == 0:
                cls = "text-center"
                val_str = str(val)
            else:
                cls = "text-left"
                val_str = str(val)
            html += f'        <td class="{cls}">{val_str}</td>\n'
        html += "      </tr>\n"

    # Строка ИТОГО
    if totals:
        html += '      <tr class="total-row">\n'
        for col_idx, val in enumerate(totals):
            if col_idx in money_cols:
                cls = "text-right"
                val_str = _format_money(val)
            elif col_idx in numeric_cols:
                cls = "text-right"
                val_str = str(val)
            else:
                cls = "text-center"
                val_str = str(val)
            html += f'        <td class="{cls}">{val_str}</td>\n'
        html += "      </tr>\n"
    html += "    </tbody>\n  </table>\n</body>\n</html>"

    with open(file_path, "w", encoding="utf-8") as f:
        f.write(html)


def _export_to_csv(file_path, title, filter_desc, now_str, headers, rows, totals):
    import csv
    with open(file_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow([title])
        writer.writerow([f"Параметры фильтрации: {filter_desc}"])
        writer.writerow([f"Дата формирования: {now_str}"])
        writer.writerow([])
        writer.writerow(headers)
        for r in rows:
            writer.writerow(r)
        if totals:
            writer.writerow(totals)


def _export_to_txt(file_path, title, filter_desc, now_str, headers, rows, totals, money_cols):
    col_widths = [len(h) for h in headers]
    for r in rows:
        for idx, val in enumerate(r):
            val_str = _format_money(val) if idx in money_cols else str(val)
            col_widths[idx] = max(col_widths[idx], len(val_str))
    if totals:
        for idx, val in enumerate(totals):
            val_str = _format_money(val) if idx in money_cols else str(val)
            col_widths[idx] = max(col_widths[idx], len(val_str))

    with open(file_path, "w", encoding="utf-8") as f:
        f.write("=" * 80 + "\n")
        f.write(f"{title}\n")
        f.write(f"Параметры фильтрации: {filter_desc}\n")
        f.write(f"Дата формирования: {now_str}\n")
        f.write("=" * 80 + "\n\n")

        # Шапка
        header_line = " | ".join(h.center(col_widths[i]) for i, h in enumerate(headers))
        f.write(header_line + "\n")
        f.write("-+-".join("-" * col_widths[i] for i in range(len(headers))) + "\n")

        # Данные
        for r in rows:
            line_parts = []
            for i, val in enumerate(r):
                val_str = _format_money(val) if i in money_cols else str(val)
                if i in money_cols or isinstance(val, (int, float)):
                    line_parts.append(val_str.rjust(col_widths[i]))
                else:
                    line_parts.append(val_str.ljust(col_widths[i]))
            f.write(" | ".join(line_parts) + "\n")

        # ИТОГО
        if totals:
            f.write("=+=".join("=" * col_widths[i] for i in range(len(headers))) + "\n")
            total_parts = []
            for i, val in enumerate(totals):
                val_str = _format_money(val) if i in money_cols else str(val)
                if i in money_cols or (val and str(val).isdigit()):
                    total_parts.append(val_str.rjust(col_widths[i]))
                else:
                    total_parts.append(val_str.center(col_widths[i]))
            f.write(" | ".join(total_parts) + "\n")
